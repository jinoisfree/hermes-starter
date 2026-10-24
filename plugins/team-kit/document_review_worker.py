"""Bounded read-only table arithmetic and document diffs. No shell or model code."""
from collections import Counter, defaultdict
import csv
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
import json
from pathlib import Path
import sys

MAX_ROWS = 20000
MAX_CELLS = 500000
MAX_ITEMS = 100


def atom(value):
    if value is None or type(value) in (str, int, float, bool):
        return value
    return str(value)


def typed(value):
    if type(value) in (int, float):
        return json.dumps(['number', str(Decimal(str(value)).normalize())])
    return json.dumps([type(value).__name__, value], ensure_ascii=False, allow_nan=False)


def table(source):
    if not isinstance(source, dict) or set(source) - {'path', 'sheet', 'header_row'}:
        raise ValueError('invalid_table_source')
    path = Path(source['path'])
    header = source.get('header_row', 1)
    if type(header) is not int or not 1 <= header <= 100:
        raise ValueError('header_row_must_be_1_to_100')
    book = stream = None
    if path.suffix.lower() == '.xlsx':
        archive_budget(path)
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheet = book[source['sheet']] if source.get('sheet') else book.worksheets[0]
        if sheet.max_column > 200 or sheet.max_row > MAX_ROWS + header:
            book.close(); raise ValueError('table_too_large_split_source_first')
        iterator = sheet.iter_rows(values_only=True)
        label = sheet.title
    else:
        stream = path.open(encoding='utf-8-sig', newline='')
        iterator = csv.reader(stream)
        label = 'CSV'
    try:
        columns, rows, cells = None, [], 0
        for number, raw in enumerate(iterator, 1):
            values = [atom(v) for v in raw]
            cells += len(values)
            if number > MAX_ROWS + header or len(values) > 200 or cells > MAX_CELLS:
                raise ValueError('table_too_large_split_source_first')
            if number < header:
                continue
            if number == header:
                # Empty trailing headings are allowed only when no data occupies them.
                while values and values[-1] in (None, ''):
                    values.pop()
                columns = [str(v).strip() if v is not None else '' for v in values]
                if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
                    raise ValueError('table_requires_unique_nonempty_column_names')
                continue
            if not any(v not in (None, '') for v in values):
                continue
            if any(v not in (None, '') for v in values[len(columns):]):
                raise ValueError('data_without_column_heading')
            values += [None] * max(0, len(columns) - len(values))
            rows.append({'row': number, 'values': dict(zip(columns, values))})
        if columns is None:
            raise ValueError('header_row_not_found')
        return {'path': str(path), 'sheet': label, 'columns': columns, 'rows': rows}
    finally:
        if book is not None: book.close()
        if stream is not None: stream.close()


def names(args, key, required=False):
    value = args.get(key, [])
    if not isinstance(value, list) or (required and not value) or len(value) > 10 or any(not isinstance(v, str) or not v for v in value) or len(set(value)) != len(value):
        raise ValueError('invalid_' + key)
    return value


def ref(t, row, column=None):
    result = {'path': t['path'], 'sheet': t['sheet'], 'row': row}
    if column:
        from openpyxl.utils import get_column_letter
        result['cell'] = get_column_letter(t['columns'].index(column) + 1) + str(row)
    return result


def analyze(args):
    sources = args.get('sources')
    if not isinstance(sources, list) or not 1 <= len(sources) <= 10:
        raise ValueError('one_to_ten_sources_required')
    tables = [table(s) for s in sources]
    if sum(len(t['rows']) * len(t['columns']) for t in tables) > MAX_CELLS:
        raise ValueError('combined_tables_too_large')
    if len({(t['path'], t['sheet']) for t in tables}) != len(tables):
        raise ValueError('duplicate_source')
    action = args['action']
    columns = names(args, 'columns', required=action == 'aggregate')
    keys = names(args, 'keys', required=action == 'compare')
    groups = names(args, 'group_by')
    for t in tables:
        if set(columns + keys + groups) - set(t['columns']):
            raise ValueError('requested_column_not_found')
    result = {'ok': True, 'sources': [{k: t[k] for k in ('path', 'sheet', 'columns')} | {'row_count': len(t['rows'])} for t in tables],
              'coverage_complete': True, 'formula_policy': 'Formula text is compared literally; never recalculated or summed.'}
    if action == 'compare':
        if len(tables) != 2 or groups:
            raise ValueError('compare_requires_two_sources_and_no_group_by')
        left, right = tables
        if not columns:
            columns = [c for c in left['columns'] if c in right['columns'] and c not in keys]
        indexes, duplicates, missing = [], [], []
        for t in tables:
            index = defaultdict(list)
            for row in t['rows']:
                key = tuple(typed(row['values'][c]) for c in keys)
                if any(row['values'][c] in (None, '') for c in keys):
                    missing.append(ref(t, row['row'])); continue
                index[key].append(row)
            for key, rows in index.items():
                if len(rows) > 1:
                    duplicates.append({'key': [json.loads(k)[1] for k in key], 'source': ref(t, rows[0]['row']), 'rows': [r['row'] for r in rows[:100]], 'row_count': len(rows)})
            indexes.append(index)
        a, b = indexes
        only_left, only_right, changes = [], [], []
        for key in dict.fromkeys([*a, *b]):
            if key not in b:
                only_left.extend(ref(left, r['row']) for r in a[key]); continue
            if key not in a:
                only_right.extend(ref(right, r['row']) for r in b[key]); continue
            if len(a[key]) != 1 or len(b[key]) != 1:
                continue  # duplicates are ambiguous, never silently choose one
            old, new = a[key][0], b[key][0]
            for c in columns:
                if typed(old['values'][c]) != typed(new['values'][c]):
                    changes.append({'key': [json.loads(k)[1] for k in key], 'column': c,
                        'before': old['values'][c], 'after': new['values'][c],
                        'before_source': ref(left, old['row'], c), 'after_source': ref(right, new['row'], c)})
        result['compared_columns'] = columns
        result['columns_only_before'] = [c for c in left['columns'] if c not in right['columns']]
        result['columns_only_after'] = [c for c in right['columns'] if c not in left['columns']]
        for name, items in [('only_before', only_left), ('only_after', only_right), ('changes', changes), ('duplicate_keys', duplicates), ('missing_keys', missing)]:
            result[name], result[name + '_count'] = items[:MAX_ITEMS], len(items)
        result['details_truncated'] = any(len(v) > MAX_ITEMS for v in (only_left, only_right, changes, duplicates, missing))
        result['comparison_complete'] = not duplicates and not missing
    elif action == 'aggregate':
        if keys:
            raise ValueError('aggregate_uses_group_by_not_keys')
        totals, excluded = {}, []
        for t in tables:
            for row in t['rows']:
                key = tuple(typed(row['values'][c]) for c in groups)
                bucket = totals.setdefault(key, {'rows': 0, 'sums': {c: Decimal(0) for c in columns}, 'numeric_counts': Counter(), 'excluded_counts': Counter()})
                bucket['rows'] += 1
                for c in columns:
                    v = row['values'][c]
                    try:
                        if v is None or isinstance(v, bool) or (isinstance(v, str) and (not v.strip() or v.startswith('='))):
                            raise InvalidOperation
                        n = Decimal(str(v))
                        if not n.is_finite(): raise InvalidOperation
                    except InvalidOperation:
                        bucket['excluded_counts'][c] += 1
                        excluded.append({'source': ref(t, row['row'], c), 'reason': 'formula' if isinstance(v, str) and v.startswith('=') else 'missing_or_non_numeric'})
                    else:
                        bucket['sums'][c] += n; bucket['numeric_counts'][c] += 1
        result.update(groups=[{'group': {c: json.loads(v)[1] for c, v in zip(groups, k)}, **b,
                              'sums': {c: str(v) if b['numeric_counts'][c] else None for c, v in b['sums'].items()}}
                             for k, b in list(totals.items())[:MAX_ITEMS]], group_count=len(totals),
                      excluded=excluded[:MAX_ITEMS], excluded_count=len(excluded), details_truncated=len(totals) > MAX_ITEMS or len(excluded) > MAX_ITEMS,
                      numeric_complete=not excluded, numeric_format='Decimal; numeric strings must not contain commas/currency symbols')
    else:
        raise ValueError('unknown_table_action')
    return result


def document(raw):
    path = Path(raw)
    items, warnings = [], []
    suffix = path.suffix.lower()
    if suffix in {'.txt', '.md'}:
        items = [(f'line:{i}', s) for i, s in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1)]
    elif suffix == '.docx':
        archive_budget(path)
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for i, block in enumerate(Document(path).iter_inner_content(), 1):
            if isinstance(block, Paragraph): items.append((f'block:{i}', block.text))
            elif isinstance(block, Table):
                for j, row in enumerate(block.rows, 1):
                    items.append((f'block:{i}/row:{j}', '\t'.join(cell.text for cell in row.cells)))
        warnings.append('Word headers, footers, comments, textboxes, nested table structure and formatting are not compared.')
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError('encrypted_or_over_200_page_pdf')
        for i, page in enumerate(reader.pages, 1):
            value = page.extract_text() or ''
            if not value.strip(): warnings.append(f'page:{i}: no text layer; image content not compared')
            items.extend((f'page:{i}/line:{j}', s) for j, s in enumerate(value.splitlines(), 1))
    elif suffix in {'.hwp', '.hwpx'}:
        import doc_worker
        # Existing readers are paginated. Fetch all pages within the report budget.
        read = doc_worker.read_hwp if suffix == '.hwp' else doc_worker.read_hwpx
        offset = 0
        while True:
            data = read({'source': str(path), 'offset': offset})
            if data['total_paragraphs'] > 10000:
                raise ValueError('document_text_too_large_split_source_first')
            items.extend((f"paragraph:{p['index']}", p['text']) for p in data['paragraphs'])
            if data['next_offset'] is None: break
            offset = data['next_offset']
        warnings.append('HWP/HWPX extracted body paragraphs only; layout and images are not compared.')
    else:
        raise ValueError('unsupported_document_type')
    if len(items) > 10000 or sum(len(v) for _, v in items) > 500000:
        raise ValueError('document_text_too_large_split_source_first')
    return items, warnings


def compare(args):
    left, lw = document(args['before']); right, rw = document(args['after'])
    changes = []
    for tag, a, b, c, d in SequenceMatcher(None, [v for _, v in left], [v for _, v in right], autojunk=False).get_opcodes():
        if tag != 'equal':
            changes.append({'change': tag, 'before': [{'location': p, 'text': v[:1000]} for p, v in left[a:b][:20]],
                            'after': [{'location': p, 'text': v[:1000]} for p, v in right[c:d][:20]],
                            'before_count': b-a, 'after_count': d-c,
                            'truncated': b-a > 20 or d-c > 20 or any(len(v) > 1000 for _, v in left[a:b] + right[c:d])})
    return {'ok': True, 'before': args['before'], 'after': args['after'], 'changes': changes[:50],
            'change_count': len(changes), 'details_truncated': len(changes) > 50 or any(x['truncated'] for x in changes),
            'body_text_equal': not changes, 'scope': 'extracted body text only; no layout, images or OCR',
            'warnings': {'before': lw, 'after': rw}, 'coverage_complete': not lw and not rw}


def archive_budget(path):
    import zipfile
    with zipfile.ZipFile(path) as archive:
        if sum(m.file_size for m in archive.infolist()) > 128 * 1024 * 1024:
            raise ValueError('expanded_document_larger_than_128MB')


if __name__ == '__main__':
    try:
        request = json.load(sys.stdin)
        result = analyze(request['args']) if request['name'] == 'table_analyze' else compare(request['args'])
        print(json.dumps(result, ensure_ascii=False, default=str, allow_nan=False))
    except Exception as exc:
        print(json.dumps({'ok': False, 'error': str(exc) if type(exc) is ValueError else type(exc).__name__}))

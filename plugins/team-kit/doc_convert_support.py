"""Local conversion helpers. No document content is logged or sent off this Mac."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unicodedata
import zipfile

OCR_MAX_PAGES = 20


def vision_binary():
    """Use the calendar helper's source-hashed, atomic Swift compilation pattern."""
    source = Path(__file__).with_name('doc_vision_ocr.swift')
    digest = hashlib.sha256(source.read_bytes()).hexdigest()[:16]
    cache = Path(os.environ.get('HERMES_HOME', str(Path.home() / '.hermes'))) / 'cache'
    if cache.is_symlink():
        raise OSError('Unsafe OCR cache')
    cache.mkdir(parents=True, exist_ok=True)
    binary = cache / ('doc-vision-' + digest)
    if binary.is_symlink():
        raise OSError('Unsafe OCR binary')
    if not binary.is_file():
        fd, temporary = tempfile.mkstemp(prefix='doc-vision-build-', dir=cache)
        os.close(fd)
        try:
            result = subprocess.run(['/usr/bin/swiftc', '-parse-as-library', str(source), '-o', temporary],
                                    capture_output=True, timeout=45)
            if result.returncode:
                raise OSError('Vision compiler unavailable')
            os.chmod(temporary, 0o700)
            os.replace(temporary, binary)
        finally:
            Path(temporary).unlink(missing_ok=True)
    return binary


def ordered_lines(lines):
    """Vision coordinates normalized to top-left. Group same-row boxes before sorting x."""
    rows = []
    for line in sorted(lines, key=lambda item: item['y']):
        if rows and abs(line['y'] - rows[-1][0]['y']) <= min(line['height'], rows[-1][0]['height']) * .5:
            rows[-1].append(line)
        else:
            rows.append([line])
    return [' '.join(item['text'] for item in sorted(row, key=lambda item: item['x'])) for row in rows]


def scanned_docx(pdf, output):
    from io import BytesIO
    from docx import Document
    from docx.shared import Inches
    if len(pdf) > OCR_MAX_PAGES:
        raise ValueError(f'스캔 PDF 글자 인식은 {OCR_MAX_PAGES}쪽까지입니다. 쪽을 나누어 변환하세요.')
    pages = None
    unavailable = False
    with tempfile.TemporaryDirectory(prefix='hermes-ocr-') as directory:
        try:
            binary = vision_binary()
            images = []
            for index, page in enumerate(pdf):
                path = Path(directory) / f'{index}.png'
                # 200 dpi for ordinary pages; bound giant page dimensions and memory.
                dpi = min(200, max(36, int(3000 * 72 / max(page.rect.width, page.rect.height))))
                page.get_pixmap(dpi=dpi).save(path)
                images.append(str(path))
            result = subprocess.run([str(binary)], input=json.dumps(images), capture_output=True,
                                    text=True, timeout=180)
            if result.returncode:
                raise OSError('Vision unavailable')
            pages = json.loads(result.stdout)
            if not isinstance(pages, list) or len(pages) != len(pdf):
                raise ValueError('Invalid OCR response')
            pages = [ordered_lines(lines) for lines in pages]
        except subprocess.TimeoutExpired:
            raise ValueError('글자 인식 시간이 초과됐습니다. 더 적은 쪽으로 나누어 변환하세요.') from None
        except (OSError, ValueError, KeyError, TypeError):
            unavailable = True
        document = Document()
        missing = []
        for index, page in enumerate(pdf):
            if index:
                document.add_page_break()
            lines = pages[index] if pages is not None and not unavailable else []
            if lines:
                for line in lines:
                    document.add_paragraph(line)
            else:
                missing.append(index + 1)
                picture = page.get_pixmap(dpi=150).tobytes('jpeg', jpg_quality=80)
                document.add_picture(BytesIO(picture), width=Inches(6))
        document.save(output)
    note = '글자 인식 결과라 오탈자가 있을 수 있고 표 구조는 문단으로 들어간다.'
    if unavailable:
        note += ' 이 환경에서 Vision 실행이 불가능해 150dpi 쪽 그림으로 저장했다.'
    elif missing:
        note += ' 글자가 인식되지 않은 쪽은 150dpi 그림으로 저장했다.'
    return [{'engine': 'PyMuPDF + python-docx' if unavailable else 'macOS Vision + python-docx',
             'pages': len(pdf), 'recognized_pages': len(pdf) - len(missing),
             'unrecognized_pages': missing, 'note': note}]


def restore_pdf_spaces(source, output):
    """Insert only source-confirmed spaces into exact character matches, preserving runs/tables.

    Ambiguous repeated passages and nonmatching paragraphs are left untouched.
    Existing whitespace and every non-whitespace character keep their positions/order.
    """
    import pymupdf
    from lxml import etree
    with pymupdf.open(source) as pdf:
        reference = '\n'.join(page.get_text() for page in pdf)
    plain, boundaries = [], set()
    whitespace = False
    for char in reference:
        if char.isspace():
            whitespace = True
        else:
            if whitespace and plain:
                boundaries.add(len(plain))
            plain.append(char)
            whitespace = False
    plain = ''.join(plain)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    inserted = 0
    cursor = 0
    for paragraph in root.findall('.//w:p', ns):
        text_nodes = paragraph.findall('.//w:t', ns)
        text = ''.join(node.text or '' for node in text_nodes)
        compact = ''.join(c for c in text if not c.isspace())
        if len(compact) < 4:
            continue
        position = plain.find(compact, cursor)
        if position < 0:
            position = plain.find(compact)
        if position < 0 or plain.find(compact, position + 1) >= 0:
            continue
        cursor = position + len(compact)
        index = 0
        previous = ''
        for node in text_nodes:
            chars = []
            for char in node.text or '':
                if not char.isspace():
                    if index and position + index in boundaries and previous and not previous.isspace():
                        chars.append(' ')
                        inserted += 1
                    index += 1
                chars.append(char)
                previous = char
            node.text = ''.join(chars)
            if node.text.startswith(' ') or node.text.endswith(' '):
                node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    if inserted:
        temporary = output.with_name(output.name + '.spacing')
        try:
            with zipfile.ZipFile(output) as original, zipfile.ZipFile(temporary, 'w') as target:
                for entry in original.infolist():
                    data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True) if entry.filename == 'word/document.xml' else original.read(entry.filename)
                    target.writestr(entry, data)
            os.replace(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
    return inserted


def normalized_font(name):
    return ''.join(name.split()).casefold()


FONT_FACES = {}


def font_inventory():
    from fontTools.ttLib import TTFont, TTCollection
    files, names, korean = [], set(), False
    FONT_FACES.clear()
    directories = (Path.home() / 'Library/Fonts', Path('/Library/Fonts'),
                   Path('/System/Library/Fonts'),
                   Path('/Applications/LibreOffice.app/Contents/Resources/fonts'))
    paths = {p.resolve() for directory in directories for p in directory.rglob('*')
             if p.suffix.lower() in {'.ttf', '.otf', '.ttc'} and p.is_file()}
    for path in sorted(paths):
        if path.suffix.lower() not in {'.ttf', '.otf', '.ttc'} or not path.is_file():
            continue
        try:
            collection = TTCollection(path, lazy=True) if path.suffix.lower() == '.ttc' else None
            fonts = collection.fonts if collection else [TTFont(path, lazy=True)]
            try:
                for font in fonts:
                    aliases = {normalized_font(record.toUnicode()) for record in font['name'].names
                               if record.nameID in {1, 4, 6, 16}}
                    faces = {normalized_font(record.toUnicode()) for record in font['name'].names if record.nameID == 6}
                    names.update(aliases)
                    for alias in aliases:
                        FONT_FACES.setdefault(alias, set()).update(faces)
                    if not korean:
                        korean = 0xAC00 in (font.getBestCmap() or {})
                files.append(path)
            finally:
                if collection:
                    collection.close()
                else:
                    fonts[0].close()
        except Exception:
            continue  # Broken user fonts must not hide other usable fonts.
    return files, names, korean


SERIF_ALIASES = ('바탕', '바탕체', '한양신명조', '한컴바탕', '궁서', '궁서체', 'Batang', 'BatangChe', 'HYSinMyeongJo', 'HCR Batang', 'Gungsuh', 'GungsuhChe')
SANS_ALIASES = ('굴림', '굴림체', '돋움', '돋움체', '맑은 고딕', '고딕', 'Gulim', 'GulimChe', 'Dotum', 'DotumChe', 'Malgun Gothic')


def requested_fonts(source):
    """Resolve fonts for live text, ignoring font-table and unused style declarations."""
    from lxml import etree
    found = set()
    with zipfile.ZipFile(source) as archive:
        def xml(path):
            return etree.fromstring(archive.read(path)) if path in archive.namelist() else None
        a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        theme = xml('word/theme/theme1.xml' if source.suffix.lower() == '.docx' else 'xl/theme/theme1.xml')
        def theme_font(group, script='latin'):
            if theme is None:
                return None
            node = theme.find('.//' + a + group + 'Font')
            if node is None:
                return None
            child = node.find(a + script)
            if child is None:
                child = next((n for n in node.findall(a + 'font') if n.get('script') == script), None)
            return child.get('typeface') if child is not None else None
        if source.suffix.lower() == '.docx':
            w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            styles = xml('word/styles.xml')
            by_id = {} if styles is None else {n.get(w+'styleId'): n for n in styles.findall(w+'style')}
            defaults = {} if styles is None else {n.get(w+'type'): n.get(w+'styleId') for n in styles.findall(w+'style') if n.get(w+'default') in {'1', 'true', 'on'}}
            def apply(properties, node):
                if node is None:
                    return
                fonts = node.find(w+'rFonts')
                if fonts is None:
                    return
                for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    # Theme takes precedence at one level; a child direct font overrides inheritance.
                    key = slot + 'Theme' if slot != 'cs' else 'cstheme'
                    direct, themed = fonts.get(w+slot), fonts.get(w+key)
                    if direct or themed:
                        properties[slot] = (direct, themed)
            def apply_style(properties, style_id, seen=None):
                seen = set() if seen is None else seen
                if style_id not in by_id or style_id in seen:
                    return
                seen.add(style_id)
                style = by_id[style_id]
                base = style.find(w+'basedOn')
                if base is not None:
                    apply_style(properties, base.get(w+'val'), seen)
                apply(properties, style.find(w+'rPr'))
            paths = ['word/document.xml'] + sorted(n for n in archive.namelist() if n.startswith(('word/header', 'word/footer')) and n.endswith('.xml'))
            for path in paths:
                root = xml(path)
                if root is None:
                    continue
                for run in root.iter(w+'r'):
                    text = ''.join(n.text or '' for n in run.iter(w+'t'))
                    if not text.strip():
                        continue
                    props = {}
                    if styles is not None:
                        apply(props, styles.find(w+'docDefaults/'+w+'rPrDefault/'+w+'rPr'))
                    paragraph = next((n for n in run.iterancestors() if n.tag == w+'p'), None)
                    table = next((n for n in run.iterancestors() if n.tag == w+'tbl'), None)
                    table_style = table.find(w+'tblPr/'+w+'tblStyle') if table is not None else None
                    if table_style is not None:
                        apply_style(props, table_style.get(w+'val'))
                    paragraph_style = paragraph.find(w+'pPr/'+w+'pStyle') if paragraph is not None else None
                    apply_style(props, paragraph_style.get(w+'val') if paragraph_style is not None else defaults.get('paragraph'))
                    rpr = run.find(w+'rPr')
                    char_style = rpr.find(w+'rStyle') if rpr is not None else None
                    apply_style(props, char_style.get(w+'val') if char_style is not None else defaults.get('character'))
                    apply(props, rpr)
                    scripts = set()
                    for char in text:
                        if char.isspace():
                            continue
                        code = ord(char)
                        if 0xAC00 <= code <= 0xD7AF or 0x1100 <= code <= 0x11FF:
                            scripts.add(('eastAsia', 'Hang'))
                        elif 0x3040 <= code <= 0x30FF:
                            scripts.add(('eastAsia', 'Jpan'))
                        elif 0x3400 <= code <= 0x9FFF:
                            scripts.add(('eastAsia', 'Hans'))
                        elif 0x590 <= code <= 0x8FF:
                            scripts.add(('cs', 'Arab' if code >= 0x600 else 'Hebr'))
                        else:
                            scripts.add(('ascii' if code < 128 else 'hAnsi', 'latin'))
                    for slot, script in scripts:
                        direct, themed = props.get(slot, (None, None))
                        name = theme_font('major' if themed and themed.startswith('major') else 'minor', script) if themed else direct
                        if name:
                            found.add(name)
        else:
            x = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            styles = xml('xl/styles.xml')
            if styles is None:
                return found
            fonts = list(styles.find(x+'fonts'))
            formats = list(styles.find(x+'cellXfs'))
            shared = xml('xl/sharedStrings.xml')
            shared = [] if shared is None else list(shared)
            def add_font(font):
                name = font.find(x+'name')
                if name is None:
                    name = font.find(x+'rFont')
                scheme = font.find(x+'scheme')
                value = theme_font(scheme.get('val')) if scheme is not None and scheme.get('val') in {'major', 'minor'} else None
                value = value or (name.get('val') if name is not None else None)
                if value:
                    found.add(value)
            for path in archive.namelist():
                if not path.startswith('xl/worksheets/') or not path.endswith('.xml'):
                    continue
                sheet = xml(path)
                cols = sheet.findall(x+'cols/'+x+'col')
                for row in sheet.findall(x+'sheetData/'+x+'row'):
                    for cell in row.findall(x+'c'):
                        if not any(cell.find(x+tag) is not None for tag in ('v', 'is', 'f')):
                            continue
                        column = 0
                        for char in cell.get('r', ''):
                            if char.isalpha():
                                column = column * 26 + ord(char.upper()) - 64
                        col_style = next((c.get('style', '0') for c in cols if int(c.get('min')) <= column <= int(c.get('max'))), '0')
                        style = int(cell.get('s', row.get('s', col_style)))
                        font = fonts[int(formats[style].get('fontId', '0'))]
                        rich = cell.find(x+'is')
                        if cell.get('t') == 's':
                            rich = shared[int(cell.findtext(x+'v'))]
                        if rich is None or rich.find(x+'t') is not None:
                            add_font(font)
                        else:
                            for run in rich.findall(x+'r'):
                                if run.findtext(x+'t', '').strip():
                                    rpr = run.find(x+'rPr')
                                    add_font(rpr if rpr is not None and rpr.find(x+'rFont') is not None else font)
    return found


def prepare_fonts(source, profile):
    from lxml import etree
    files, installed, korean = font_inventory()
    if not korean:
        raise ValueError('한글 글꼴이 없습니다. brew install --cask font-nanum-gothic font-nanum-myeongjo 로 설치하세요.')
    font_dir = profile / 'user/fonts'
    font_dir.mkdir(parents=True)
    for index, font in enumerate(files):
        (font_dir / f'{index}-{font.name}').symlink_to(font)
    requested = requested_fonts(source)
    known = {normalized_font(name): 'NanumMyeongjo' for name in SERIF_ALIASES}
    known.update({normalized_font(name): 'NanumGothic' for name in SANS_ALIASES})
    known.update({'calibri': 'Carlito', 'cambria': 'Caladea'})
    substitutes = []
    automatic = []
    for name in sorted(requested):
        if normalized_font(name) in installed:
            continue
        target = known.get(normalized_font(name))
        if target and normalized_font(target) in installed:
            substitutes.append({'from': name, 'to': target})
        else:
            automatic.append(name)
    oor = 'http://openoffice.org/2001/registry'
    root = etree.Element(f'{{{oor}}}items', nsmap={'oor': oor})
    item = etree.SubElement(root, 'item', {f'{{{oor}}}path': '/org.openoffice.Office.Common/Font/Substitution'})
    prop = etree.SubElement(item, 'prop', {f'{{{oor}}}name': 'Replacement', f'{{{oor}}}op': 'fuse'})
    etree.SubElement(prop, 'value').text = 'true'
    pairs = etree.SubElement(root, 'item', {f'{{{oor}}}path': '/org.openoffice.Office.Common/Font/Substitution/FontPairs'})
    # Office imports may canonicalize Korean family names (바탕 → Batang).
    # Register every alias, but never override an actually installed family.
    replacements = [{'from': name, 'to': target} for name, target in
                    [(name, 'NanumMyeongjo') for name in SERIF_ALIASES] +
                    [(name, 'NanumGothic') for name in SANS_ALIASES] +
                    [('Calibri', 'Carlito'), ('Cambria', 'Caladea')]
                    if normalized_font(name) not in installed and normalized_font(target) in installed]
    for index, pair in enumerate(replacements):
        node = etree.SubElement(pairs, 'node', {f'{{{oor}}}name': f'_{index}', f'{{{oor}}}op': 'replace'})
        for key, value in {'ReplaceFont': pair['from'], 'SubstituteFont': pair['to'], 'Always': 'true', 'OnScreenOnly': 'false'}.items():
            prop = etree.SubElement(node, 'prop', {f'{{{oor}}}name': key})
            etree.SubElement(prop, 'value').text = value
    (profile / 'user/registrymodifications.xcu').write_bytes(etree.tostring(root, xml_declaration=True, encoding='UTF-8'))
    return substitutes, automatic


def verify_font_report(source, output, substitutes, automatic):
    """Do not claim a requested installed/replacement family was used without PDF evidence.

    FONT_FACES is populated by prepare_fonts in this worker process. Match PostScript
    names (including bold/italic faces), after removing PDF subset prefixes.
    """
    import pymupdf
    with pymupdf.open(output) as pdf:
        embedded = {normalized_font(font[3].split('+')[-1]) for page in pdf for font in page.get_fonts()}
    def used(name):
        return bool(FONT_FACES.get(normalized_font(name), set()) & embedded)
    verified = [pair for pair in substitutes if used(pair['to'])]
    missing = set(automatic)
    missing.update(pair['from'] for pair in substitutes if not used(pair['to']))
    replaced = {pair['from'] for pair in substitutes}
    for name in requested_fonts(source):
        if name not in replaced and normalized_font(name) in FONT_FACES and not used(name):
            missing.add(name)
    return verified, sorted(missing)


def prepare_xlsx(source, destination, fit='width', widen_columns=False):
    """Edit only worksheet XML in a temporary ZIP copy; drawings/charts stay byte-identical."""
    from copy import deepcopy
    from lxml import etree
    import re
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    q = lambda name: '{' + ns + '}' + name
    def insert_before(root, node, names):
        for index, child in enumerate(root):
            if etree.QName(child).localname in names:
                root.insert(index, node)
                return
        root.append(node)
    widened = kept = overridden_scale = 0
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, 'w') as output:
        strings = []
        if widen_columns and 'xl/sharedStrings.xml' in original.namelist():
            strings = [''.join(node.itertext()) for node in etree.fromstring(original.read('xl/sharedStrings.xml'))]
        for entry in original.infolist():
            data = original.read(entry.filename)
            if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', entry.filename) and (fit != 'none' or widen_columns):
                root = etree.fromstring(data)
                existing_pr = root.find(q('sheetPr') + '/' + q('pageSetUpPr'))
                preserve = fit == 'width' and existing_pr is not None and existing_pr.get('fitToPage') in {'1', 'true', 'on'}
                if preserve:
                    kept += 1
                if fit != 'none' and not preserve:
                    sheet_pr = root.find(q('sheetPr'))
                    if sheet_pr is None:
                        sheet_pr = etree.Element(q('sheetPr')); root.insert(0, sheet_pr)
                    setup_pr = sheet_pr.find(q('pageSetUpPr'))
                    if setup_pr is None:
                        setup_pr = etree.SubElement(sheet_pr, q('pageSetUpPr'))
                    setup_pr.set('fitToPage', '1')
                    setup = root.find(q('pageSetup'))
                    if setup is None:
                        setup = etree.Element(q('pageSetup'))
                        insert_before(root, setup, {'headerFooter', 'rowBreaks', 'colBreaks', 'customProperties', 'cellWatches', 'ignoredErrors', 'smartTags', 'drawing', 'legacyDrawing', 'legacyDrawingHF', 'picture', 'oleObjects', 'controls', 'webPublishItems', 'tableParts', 'extLst'})
                    if fit == 'width' and 'scale' in setup.attrib:
                        overridden_scale += 1
                    setup.attrib.pop('scale', None)
                    setup.set('fitToWidth', '1'); setup.set('fitToHeight', '1' if fit == 'sheet' else '0')
                if widen_columns:
                    widths = {}
                    merged = {n.get('ref').split(':')[0] for n in root.findall('./'+q('mergeCells')+'/'+q('mergeCell'))}
                    for cell in root.findall('./'+q('sheetData')+'/'+q('row')+'/'+q('c')):
                        address = cell.get('r', '')
                        if address in merged:
                            continue
                        letters = ''.join(c for c in address if c.isalpha())
                        col = 0
                        for char in letters:
                            col = col * 26 + ord(char.upper()) - 64
                        value = cell.findtext(q('v'), '')
                        if cell.get('t') == 's':
                            value = strings[int(value)]
                        elif cell.get('t') == 'inlineStr':
                            value = ''.join(cell.find(q('is')).itertext())
                        # Excel width units approximate the default font's digit width.
                        width = max((sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in line) for line in value.splitlines()), default=0)
                        if col and width:
                            widths[col] = max(widths.get(col, 0), min(255, width * 1.15 + 3))
                    cols = root.find(q('cols'))
                    if cols is None:
                        cols = etree.Element(q('cols')); insert_before(root, cols, {'sheetData'})
                    existing = list(cols)
                    for node in existing:
                        low, high = int(node.get('min')), int(node.get('max'))
                        changed = sorted(c for c in widths if low <= c <= high and widths[c] > float(node.get('width', '8.43')))
                        if not changed:
                            continue
                        cols.remove(node)
                        start = low
                        for col in changed:
                            if start < col:
                                part = deepcopy(node); part.set('min', str(start)); part.set('max', str(col-1)); cols.append(part)
                            part = deepcopy(node); part.set('min', str(col)); part.set('max', str(col)); part.set('width', str(widths[col])); part.set('customWidth', '1'); cols.append(part); widened += 1
                            start = col + 1
                        if start <= high:
                            part = deepcopy(node); part.set('min', str(start)); cols.append(part)
                    for col, width in widths.items():
                        if any(int(n.get('min')) <= col <= int(n.get('max')) for n in existing):
                            continue
                        if width > 8.43:
                            etree.SubElement(cols, q('col'), min=str(col), max=str(col), width=str(width), customWidth='1'); widened += 1
                    cols[:] = sorted(cols, key=lambda n: int(n.get('min')))
                data = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            output.writestr(entry, data)
    return {'widened_columns': widened, 'kept_print_setup': kept,
            'overridden_scale_setup': overridden_scale}

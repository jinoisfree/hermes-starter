"""Read, edit or render one document. Runs under ~/.local/share/hermes-docs, not the gateway.

Reads one JSON job on stdin and prints one JSON result. An edit only ever writes ``job["output"]``,
a temporary file the caller moves into place after checking the result; a render writes page images
into ``job["output_dir"]``; a read writes nothing. A fill (doc_fill) writes one file per data row
into ``job["staging"]`` with fill_template_engine.py; the caller links them into the output folder.
"""
from __future__ import annotations

import json
import sys


def _sheet(book, name):
    if name in (None, ""):
        return book.active
    if name not in book.sheetnames:
        raise ValueError(f"no sheet named {name!r}; sheets: {', '.join(book.sheetnames)}")
    return book[name]


def xlsx(job):
    import openpyxl
    book = openpyxl.load_workbook(job["source"]) if job.get("source") else openpyxl.Workbook()
    done = []
    for op in job["operations"]:
        kind = op.get("op")
        if kind == "set_cells":
            sheet = _sheet(book, op.get("sheet"))
            cells = op.get("cells")
            if not isinstance(cells, dict) or not cells:
                raise ValueError("set_cells needs cells: {\"A1\": value}")
            for ref, value in cells.items():
                sheet[str(ref)] = value
            done.append(f"set_cells {sheet.title}: {len(cells)}")
        elif kind == "append_rows":
            sheet = _sheet(book, op.get("sheet"))
            rows = op.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) for row in rows):
                raise ValueError("append_rows needs rows: a list of lists")
            for row in rows:
                sheet.append(row)
            done.append(f"append_rows {sheet.title}: {len(rows)}")
        elif kind == "delete_rows":
            sheet = _sheet(book, op.get("sheet"))
            start, count = int(op.get("start", 0)), int(op.get("count", 1))
            if start < 1 or count < 1 or start + count - 1 > sheet.max_row:
                raise ValueError(f"delete_rows: rows {start}..{start + count - 1} are outside 1..{sheet.max_row}")
            sheet.delete_rows(start, count)
            done.append(f"delete_rows {sheet.title}: {start}+{count}")
        elif kind == "insert_rows":
            sheet = _sheet(book, op.get("sheet"))
            start, count = int(op.get("start", 0)), int(op.get("count", 1))
            if start < 1 or not 1 <= count <= 1000:
                raise ValueError("insert_rows needs start >= 1 and count 1..1000")
            sheet.insert_rows(start, count)
            done.append(f"insert_rows {sheet.title}: {start}+{count}")
        elif kind == "replace_text":
            find, replace = op.get("find"), op.get("replace")
            if not isinstance(find, str) or not find or not isinstance(replace, str):
                raise ValueError("replace_text needs find and replace strings")
            sheets = [_sheet(book, op["sheet"])] if op.get("sheet") else book.worksheets
            count = 0
            for sheet in sheets:
                for cells in sheet.iter_rows():
                    for cell in cells:
                        if isinstance(cell.value, str) and find in cell.value and not cell.value.startswith("="):
                            count += cell.value.count(find)
                            cell.value = cell.value.replace(find, replace)
            if not count:
                raise ValueError(f"replace_text: {find!r} was not found in any text cell")
            done.append(f"replace_text: {count}")
        elif kind == "add_sheet":
            name = str(op.get("name") or "").strip()
            if not name or name in book.sheetnames:
                raise ValueError("add_sheet needs a new sheet name")
            book.create_sheet(name)
            done.append(f"add_sheet {name}")
        elif kind == "rename_sheet":
            sheet = _sheet(book, op.get("sheet"))
            name = str(op.get("name") or "").strip()
            if not name or name in book.sheetnames:
                raise ValueError("rename_sheet needs a new sheet name")
            done.append(f"rename_sheet {sheet.title} → {name}")
            sheet.title = name
        elif kind == "delete_sheet":
            sheet = _sheet(book, op.get("sheet"))
            if len(book.sheetnames) == 1:
                raise ValueError("delete_sheet: a workbook keeps at least one sheet")
            done.append(f"delete_sheet {sheet.title}")
            book.remove(sheet)
        else:
            raise ValueError(f"unknown Excel operation {kind!r}")
    book.save(job["output"])
    return done


def _replace_in_paragraph(paragraph, find, replace):
    count = paragraph.text.count(find)
    if not count:
        return 0
    for run in paragraph.runs:  # keeps formatting when the text sits in one run
        if find in run.text:
            run.text = run.text.replace(find, replace)
    if find in paragraph.text and paragraph.runs:  # split across runs: keep the first run's formatting
        text = paragraph.text.replace(find, replace)
        for run in paragraph.runs[1:]:
            run.text = ""
        paragraph.runs[0].text = text
    return count


def docx(job):
    import docx as word
    document = word.Document(job["source"]) if job.get("source") else word.Document()
    done = []
    for op in job["operations"]:
        kind = op.get("op")
        if kind == "replace_text":
            find, replace = op.get("find"), op.get("replace")
            if not isinstance(find, str) or not find or not isinstance(replace, str):
                raise ValueError("replace_text needs find and replace strings")
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            count = sum(_replace_in_paragraph(paragraph, find, replace) for paragraph in paragraphs)
            if not count:
                raise ValueError(f"replace_text: {find!r} was not found")
            done.append(f"replace_text: {count}")
        elif kind in ("set_paragraph", "delete_paragraph", "insert_paragraph"):
            paragraphs = document.paragraphs
            index = op.get("index")
            if not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < len(paragraphs):
                raise ValueError(f"{kind}: index must be 0..{len(paragraphs) - 1} (see doc_read)")
            target = paragraphs[index]
            if kind == "delete_paragraph":
                target._element.getparent().remove(target._element)
            elif kind == "set_paragraph":
                for run in target.runs[1:]:
                    run.text = ""
                if target.runs:
                    target.runs[0].text = str(op.get("text", ""))
                else:
                    target.add_run(str(op.get("text", "")))
            else:
                try:
                    target.insert_paragraph_before(str(op.get("text", "")), style=op.get("style") or None)
                except KeyError:
                    raise ValueError(f"insert_paragraph: this document has no style {op.get('style')!r}")
            done.append(f"{kind}: {index}")
        elif kind == "set_table_cell":
            try:
                cell = document.tables[int(op["table"])].cell(int(op["row"]), int(op["col"]))
            except (KeyError, IndexError, ValueError, TypeError):
                raise ValueError("set_table_cell needs table, row and col inside the document (see doc_read)")
            cell.text = str(op.get("text", ""))
            done.append(f"set_table_cell: {op['table']}/{op['row']}/{op['col']}")
        elif kind == "append_paragraphs":
            items = op.get("paragraphs")
            if not isinstance(items, list) or not items:
                raise ValueError("append_paragraphs needs paragraphs: [{\"text\": ..., \"style\": optional}]")
            for item in items:
                item = item if isinstance(item, dict) else {"text": str(item)}
                style = item.get("style")
                try:
                    document.add_paragraph(str(item.get("text", "")), style=style or None)
                except KeyError:
                    raise ValueError(f"append_paragraphs: this document has no style {style!r}")
            done.append(f"append_paragraphs: {len(items)}")
        elif kind == "append_table":
            rows = op.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) and row for row in rows):
                raise ValueError("append_table needs rows: a list of non-empty lists")
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = "" if value is None else str(value)
            done.append(f"append_table: {len(rows)}x{width}")
        else:
            raise ValueError(f"unknown Word operation {kind!r}")
    document.save(job["output"])
    return done


def _pages(spec, total):
    if spec in (None, ""):
        return list(range(total))
    picked = []
    for part in str(spec).replace(" ", "").split(","):
        first, _, last = part.partition("-")
        start, end = int(first), int(last or first)
        if start < 1 or end < start or end > total:
            raise ValueError(f"pages {part!r} are outside 1..{total}")
        picked.extend(range(start - 1, end))
    return picked


def pdf(job):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    done = []
    for source in job["sources"]:
        reader = PdfReader(source["path"])
        if reader.is_encrypted:
            raise ValueError(f"{source['path']} is password-protected")
        rotate = int(source.get("rotate") or 0)
        if rotate % 90:
            raise ValueError("rotate must be a multiple of 90")
        picked = _pages(source.get("pages"), len(reader.pages))
        for index in picked:
            page = reader.pages[index]
            if rotate:
                page.rotate(rotate)
            writer.add_page(page)
        done.append(f"{source['path']}: {len(picked)} page(s)")
    with open(job["output"], "wb") as handle:
        writer.write(handle)
    return done


def _clip(value, limit=500):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    return text if len(text) <= limit else text[:limit] + "…"


def read_pdf(job):
    """Read the local text layer with explicit page coverage; never silently skip scans."""
    from pathlib import Path
    from pypdf import PdfReader
    if Path(job["source"]).stat().st_size > 50 * 1024 * 1024:
        raise ValueError("PDF 읽기는 파일당 50MB까지입니다.")
    reader = PdfReader(job["source"])
    if reader.is_encrypted:
        raise ValueError("암호 PDF는 읽을 수 없습니다. 암호를 해제한 사본을 요청하세요.")
    total = len(reader.pages)
    picked = list(dict.fromkeys(_pages(job["pages"], total))) if job.get("pages") else list(range(min(total, 5)))
    if len(picked) > 5:
        raise ValueError("PDF는 한 번에 5쪽까지 읽습니다. pages로 범위를 나누세요.")
    shown, empty = [], []
    for index in picked:
        text = reader.pages[index].extract_text() or ""
        if not text.strip():
            empty.append(index + 1)
        shown.append({"page": index + 1, "text": text[:20000], "truncated": len(text) > 20000})
    next_page = max(picked) + 2 if picked and max(picked) + 1 < total else None
    return {"pages": shown, "total_pages": total, "next_page": next_page,
            "textless_pages": empty, "complete": len(picked) == total and not empty and not any(p["truncated"] for p in shown),
            "note": "글자 층만 읽었다. 사진·서명·도표와 글자 없는 쪽은 pdf_render → vision_analyze로 확인한다. 잘린 쪽도 같은 방식으로 확인한다."}


def read_xlsx(job):
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries
    book = openpyxl.load_workbook(job["source"], read_only=False, data_only=False)
    sheets = [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column} for sheet in book.worksheets]
    sheet = _sheet(book, job.get("sheet"))
    if job.get("range"):
        left, top, right, bottom = range_boundaries(str(job["range"]))
    else:
        left, top, right, bottom = 1, 1, min(sheet.max_column, 30), min(sheet.max_row, 200)
    right, bottom = min(right, left + 49), min(bottom, top + 499)
    rows = []
    for number in range(top, bottom + 1):
        cells = {}
        for column in range(left, right + 1):
            value = sheet.cell(row=number, column=column).value
            if value not in (None, ""):
                cells[f"{get_column_letter(column)}{number}"] = _clip(value)
        if cells:
            rows.append(cells)
    return {"sheets": sheets, "sheet": sheet.title, "shown": f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
            "cells": rows, "merged": [str(item) for item in list(sheet.merged_cells.ranges)[:50]],
            "note": "values starting with '=' are formulas; empty cells are omitted"}


def read_docx(job):
    import docx as word
    document = word.Document(job["source"])
    paragraphs = document.paragraphs
    start = max(int(job.get("offset") or 0), 0)
    shown = [{"index": index, "style": paragraph.style.name if paragraph.style is not None else None, "text": _clip(paragraph.text, 1000)}
             for index, paragraph in enumerate(paragraphs[start:start + 200], start)]
    tables = []
    for number, table in enumerate(document.tables[:20]):
        tables.append({"table": number, "rows": [[_clip(cell.text, 200) for cell in row.cells] for row in table.rows[:50]],
                       "total_rows": len(table.rows)})
    return {"paragraphs": shown, "total_paragraphs": len(paragraphs), "tables": tables, "total_tables": len(document.tables)}


# Bound both archive members and decompressed OLE streams before parsing.
MAX_HANGUL_BYTES = 32 * 1024 * 1024


def _paragraphs(texts, job):
    start = max(int(job.get("offset") or 0), 0)
    return {"paragraphs": [{"index": i, "text": text} for i, text in
                           enumerate(texts[start:start + 200], start)],
            "total_paragraphs": len(texts),
            "next_offset": start + 200 if start + 200 < len(texts) else None}


def read_hwpx(job):
    import re
    import zipfile
    import xml.etree.ElementTree as ET
    try:
        with zipfile.ZipFile(job["source"]) as archive:
            members = archive.infolist()
            if any(m.flag_bits & 1 for m in members):
                raise ValueError("암호 한글 문서는 읽을 수 없습니다. PDF를 요청하세요.")
            sections = sorted((m for m in members if re.fullmatch(r"Contents/section[0-9]+\.xml", m.filename)),
                              key=lambda m: int(re.search(r"section([0-9]+)", m.filename)[1]))
            if not sections or sum(m.file_size for m in sections) > MAX_HANGUL_BYTES:
                raise ValueError("HWPX 본문이 없거나 읽기 크기 한도를 넘었습니다.")
            texts = []
            def walk(node):
                # Flush around nested table paragraphs: don't repeat their text in the parent.
                if node.tag.rsplit("}", 1)[-1] == "p":
                    parts = []
                    def content(child):
                        local = child.tag.rsplit("}", 1)[-1]
                        if local == "p":
                            if parts:
                                texts.append("".join(parts)); parts.clear()
                            walk(child)
                        elif local == "t":
                            parts.append(child.text or "")
                            for sub in child:
                                if sub.tag.rsplit("}", 1)[-1] in {"tab", "lineBreak"}:
                                    parts.append("\t" if sub.tag.endswith("tab") else "\n")
                                parts.append(sub.tail or "")
                        else:
                            for sub in child:
                                content(sub)
                    for child in node:
                        content(child)
                    if parts:
                        texts.append("".join(parts))
                else:
                    for child in node:
                        walk(child)
            for member in sections:
                raw = archive.read(member)
                if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                    raise ValueError("지원하지 않는 HWPX XML 선언입니다.")
                walk(ET.fromstring(raw))
            return _paragraphs(texts, job)
    except (OSError, zipfile.BadZipFile, ET.ParseError, RuntimeError) as exc:
        raise ValueError("손상되었거나 암호화된 HWPX 문서를 읽지 못했습니다.") from exc


def _hwp_text(raw):
    import struct
    if len(raw) % 2:
        raise ValueError("손상된 HWP 문단입니다.")
    units = struct.unpack('<' + 'H' * (len(raw) // 2), raw)
    out = bytearray()
    i = 0
    # HWP5 inline/extended controls occupy eight UTF-16 code units.
    long_controls = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}
    while i < len(units):
        char = units[i]
        if char in long_controls:
            if i + 8 > len(units):
                raise ValueError("손상된 HWP 제어 문자입니다.")
            if char == 9:
                out.extend(b'\t\x00')
            i += 8
        else:
            if char >= 32 or char == 10:
                out.extend(struct.pack('<H', char))
            i += 1
    return out.decode('utf-16le')


def read_hwp(job):
    import re
    import struct
    import zlib
    import olefile
    try:
        with olefile.OleFileIO(job["source"]) as document:
            header = document.openstream('FileHeader').read()
            if len(header) < 40 or not header.startswith(b'HWP Document File') or header[35] != 5:
                raise ValueError("지원하지 않거나 손상된 HWP5 문서입니다.")
            flags = struct.unpack_from('<I', header, 36)[0]
            if flags & 6 or document.exists('ViewText'):
                raise ValueError("암호·배포용 한글 문서는 읽을 수 없습니다. PDF를 요청하세요.")
            sections = sorted((p for p in document.listdir() if len(p) == 2 and p[0] == 'BodyText'
                               and re.fullmatch(r'Section[0-9]+', p[1])), key=lambda p: int(p[1][7:]))
            if not sections:
                raise ValueError("HWP 본문을 찾을 수 없습니다.")
            texts, total = [], 0
            for section in sections:
                if document.get_size(section) > MAX_HANGUL_BYTES:
                    raise ValueError("HWP 읽기 크기 한도를 넘었습니다.")
                raw = document.openstream(section).read()
                if flags & 1:
                    decoder = zlib.decompressobj(-15)
                    raw = decoder.decompress(raw, MAX_HANGUL_BYTES + 1)
                    trailer = decoder.unused_data
                    trailer_ok = not trailer or (len(trailer) == 8 and
                        struct.unpack('<II', trailer) == (zlib.crc32(raw), len(raw)))
                    if not decoder.eof or not trailer_ok:
                        raise ValueError("손상되었거나 너무 큰 HWP 압축 본문입니다.")
                total += len(raw)
                if total > MAX_HANGUL_BYTES:
                    raise ValueError("HWP 읽기 크기 한도를 넘었습니다.")
                offset = 0
                while offset < len(raw):
                    if offset + 4 > len(raw):
                        raise ValueError("손상된 HWP 레코드입니다.")
                    record = struct.unpack_from('<I', raw, offset)[0]; offset += 4
                    size = record >> 20
                    if size == 0xfff:
                        if offset + 4 > len(raw):
                            raise ValueError("손상된 HWP 레코드 길이입니다.")
                        size = struct.unpack_from('<I', raw, offset)[0]; offset += 4
                    if offset + size > len(raw):
                        raise ValueError("잘린 HWP 레코드입니다.")
                    if record & 0x3ff == 67:
                        texts.append(_hwp_text(raw[offset:offset + size]))
                    offset += size
            return _paragraphs(texts, job)
    except (OSError, zlib.error, UnicodeError, struct.error) as exc:
        raise ValueError("손상된 HWP 문서를 읽지 못했습니다.") from exc


def render(job):
    import os
    import pypdfium2
    document = pypdfium2.PdfDocument(job["source"])
    total = len(document)
    picked = _pages(job.get("pages") or "1", total)
    if len(picked) > job.get("limit", 5):
        raise ValueError(f"at most {job.get('limit', 5)} pages per call; this PDF has {total}")
    os.makedirs(job["output_dir"], exist_ok=True)
    images = []
    for index in picked:
        page = document[index]
        scale = min(2.5, 1800 / max(page.get_width(), 1))
        path = os.path.join(job["output_dir"], f"page-{index + 1}.png")
        page.render(scale=scale).to_pil().save(path)
        images.append({"page": index + 1, "image": path})
    return {"total_pages": total, "images": images}


def _conversion_support():
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location('doc_convert_support', Path(__file__).with_name('doc_convert_support.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def convert(job):
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    support = _conversion_support()
    source = Path(job["source"])
    output = Path(job["output"])
    target_format = job["format"]
    pair = (source.suffix.lower(), target_format)
    if pair not in {(".pdf", "docx"), (".docx", "pdf"), (".xlsx", "pdf")}:
        raise ValueError("supports only .pdf → .docx, .docx → .pdf, .xlsx → .pdf")
    if any(key in job for key in ('fit', 'widen_columns')) and source.suffix.lower() != '.xlsx':
        raise ValueError('fit and widen_columns are only supported for .xlsx sources')
    if job.get('fit', 'width') not in ('width', 'sheet', 'none') or type(job.get('widen_columns', False)) is not bool:
        raise ValueError('fit must be width/sheet/none; widen_columns must be boolean')
    if source.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("source exceeds 100 MB")
    if pair == (".pdf", "docx"):
        import pymupdf
        from pdf2docx import Converter
        from pypdf import PdfReader
        if PdfReader(source).is_encrypted:
            raise ValueError("encrypted PDFs are not supported")
        with pymupdf.open(source) as pdf:
            if pdf.is_encrypted:
                raise ValueError("encrypted PDFs are not supported")
            if not 1 <= len(pdf) <= 300:
                raise ValueError("PDF must have 1 to 300 pages")
            image_pages = [index + 1 for index, page in enumerate(pdf) if not page.get_text().strip()]
            if len(image_pages) == len(pdf):
                return support.scanned_docx(pdf, output)
        converter = Converter(str(source))
        try:
            converter.convert(str(output), multi_processing=False)
        finally:
            converter.close()
        spaces = support.restore_pdf_spaces(source, output)
        note = "표·배치·글꼴이 원본과 다를 수 있으므로 결과를 확인한다."
        if image_pages:
            note += " 글자 층 없는 쪽은 OCR하지 않았다: " + ', '.join(map(str, image_pages))
        return [{"engine": "pdf2docx", "image_only_pages": image_pages, "restored_spaces": spaces, "note": note}]

    soffice = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    executable = str(soffice) if soffice.is_file() else shutil.which("soffice")
    if not executable:
        raise ValueError("LibreOffice is missing; install with brew install --cask libreoffice")
    with tempfile.TemporaryDirectory(prefix="hermes-convert-") as folder:
        work = Path(folder)
        substitutions, automatic = support.prepare_fonts(source, work / 'profile')
        local_source = work / ("input" + source.suffix.lower())
        print_setup = {}
        if source.suffix.lower() == '.xlsx':
            print_setup = support.prepare_xlsx(source, local_source, job.get('fit', 'width'), job.get('widen_columns', False))
        else:
            shutil.copyfile(source, local_source)
        result = subprocess.run([executable, "-env:UserInstallation=" + (work / "profile").as_uri(),
                                 "--headless", "--convert-to", "pdf", "--outdir", str(work), str(local_source)],
                                capture_output=True, text=True, timeout=270)
        converted = work / "input.pdf"
        if result.returncode != 0 or not converted.is_file() or not converted.stat().st_size:
            raise ValueError("LibreOffice conversion failed: " + (result.stderr or result.stdout)[-300:])
        substitutions, automatic = support.verify_font_report(source, converted, substitutions, automatic)
        shutil.move(str(converted), output)
    return [{"engine": "LibreOffice", "font_substitutions": substitutions,
             "unmapped_missing_fonts": automatic,
             **({"fit": job.get('fit', 'width'), **print_setup} if source.suffix.lower() == '.xlsx' else {}),
             "note": "글꼴 대체로 줄바꿈이 달라질 수 있다. unmapped_missing_fonts는 미설치 또는 결과 PDF에서 요청 글꼴을 확인하지 못한 이름이며 대체 글꼴을 단정하지 않는다. Excel 인쇄 영역·배율·열 너비를 결과에서 확인한다."}]


FILL_FIELD = r"\{([^{}]+)\}"
FILL_NAME_CHARS = 80


def _fill_value(value):
    """Values go in as stored: dates as YYYY-MM-DD, whole floats without .0, text untouched."""
    import datetime
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d" if value.time() == datetime.time() else "%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fill_engine():
    import fill_template_engine as engine
    engine._missing = lambda name: f"«누락: {name}»"  # visible in the file, never a silent blank
    engine._fmt_value = _fill_value
    return engine


def _fill_plan(engine, job):
    """Tokens, columns, their mapping and the selected rows; raises on a bad mapping or pattern."""
    import re
    tokens = engine.tokens_in(job["template"])
    plan = {"tokens": tokens}
    if not job.get("data"):
        return plan, []
    headers, rows = engine.load_rows(job["data"])
    headers = [h for h in headers if h]
    rows = [{str(k).strip(): v for k, v in row.items() if k is not None} for row in rows]
    lower = {h.lower(): h for h in headers}
    columns = {}
    for token, column in (job.get("columns") or {}).items():
        if str(column).strip().lower() not in lower:
            raise ValueError(f"columns: no column {column!r} in the data; columns: {', '.join(headers)}")
        columns[token] = lower[str(column).strip().lower()]
    mapping = {t: c for t, c in engine._resolve_map(tokens, columns, headers).items() if t in tokens}
    fields = re.findall(FILL_FIELD, job.get("name_pattern") or "")
    for field in fields:
        if field.strip().lower() not in lower:
            raise ValueError(f"name_pattern: no column {field!r} in the data; columns: {', '.join(headers)}")
    start = job.get("start") or 1
    selected = rows[start - 1:start - 1 + job["limit"]] if job.get("limit") else rows[start - 1:]
    blanks = {t: sum(1 for row in rows if _fill_value(row.get(c)) == "") for t, c in mapping.items()}
    plan.update(columns=headers, mapping=mapping, unmapped_tokens=[t for t in tokens if t not in mapping],
                unused_columns=[h for h in headers if h not in mapping.values()], rows=len(rows),
                start=start, selected=len(selected), blank_values={t: n for t, n in blanks.items() if n})
    return plan, selected


def _fill_names(engine, job, rows, offset, ext):
    import re
    from pathlib import Path
    pattern, used, names = job.get("name_pattern"), set(), []
    default = Path(job["template"]).stem
    for n, row in enumerate(rows, start=offset):
        lower = {k.lower(): v for k, v in row.items()}
        stem = (re.sub(FILL_FIELD, lambda m: _fill_value(lower.get(m.group(1).strip().lower())) or f"{n}행", pattern)
                if pattern else f"{default}_{n}")
        stem = engine._safe_name(stem)[:FILL_NAME_CHARS].strip() or f"{n}행"
        name, k = stem, 2
        while name.lower() in used:  # macOS names are case-insensitive
            name, k = f"{stem} ({k})", k + 1
        used.add(name.lower())
        names.append(name + ext)
    return names


def fill_inspect(job):
    engine = _fill_engine()
    plan, rows = _fill_plan(engine, job)
    if rows:
        plan["sample_names"] = _fill_names(engine, job, rows[:3], plan["start"], "." + job["template"].rsplit(".", 1)[-1].lower())
        plan["sample"] = {t: _fill_value(rows[0].get(c)) for t, c in plan["mapping"].items()}
    return plan


def fill(job):
    """One file per selected row into job["staging"], an empty folder the caller links files out of."""
    from pathlib import Path
    engine = _fill_engine()
    plan, rows = _fill_plan(engine, job)
    if not plan["tokens"]:
        raise ValueError("the template has no {{tokens}}")
    if not rows:
        raise ValueError("no data rows selected")
    if len(rows) > job["max_rows"]:
        raise ValueError(f"{len(rows)} rows selected; at most {job['max_rows']} per call (use start and limit)")
    ext = Path(job["template"]).suffix.lower()
    staging, written = Path(job["staging"]), []
    for n, (row, name) in enumerate(zip(rows, _fill_names(engine, job, rows, plan["start"], ext)), start=plan["start"]):
        values = {t: _fill_value(row.get(c)) for t, c in plan["mapping"].items()}
        missing = engine._fill_one(job["template"], values, str(staging / name))
        written.append({"name": name, "row": n, "missing": missing})
    return {**plan, "files": written}


def main():
    try:
        job = json.loads(sys.stdin.read())
        if job["kind"] in ("read_pdf", "read_xlsx", "read_docx", "read_hwp", "read_hwpx", "render", "fill_inspect", "fill"):
            result = {"read_pdf": read_pdf, "read_xlsx": read_xlsx, "read_docx": read_docx, "read_hwp": read_hwp, "read_hwpx": read_hwpx, "render": render,
                      "fill_inspect": fill_inspect, "fill": fill}[job["kind"]](job)
            print(json.dumps({"ok": True, **result}, ensure_ascii=False, default=str))
            return
        done = {"xlsx": xlsx, "docx": docx, "pdf": pdf, "convert": convert}[job["kind"]](job)
        print(json.dumps({"ok": True, "done": done}, ensure_ascii=False))
    except Exception as exc:  # the caller reports it and discards the temporary file
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"[:500]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
